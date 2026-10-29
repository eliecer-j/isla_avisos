"""Genera un Excel de pruebas (tareas_prueba.xlsx) con fechas relativas a hoy."""
from __future__ import annotations

import sys
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = ["Fecha", "Hora", "Tarea", "Detalle", "Notas", "Prioridad", "Estado"]


def build_rows() -> list:
    today = date.today()
    return [
        (today + timedelta(days=2), None, "Juan Pérez",
         "Contrato temporal", "Renovar o dar de baja", "Crítica", ""),
        (today + timedelta(days=5), time(9, 0), "María Gómez",
         "Contrato de prácticas", "Informe de tutor", "Alta", ""),
        (today + timedelta(days=15), None, "Carlos Ruiz",
         "Contrato por obra", "Revisar con el jefe de área", "Alta", ""),
        (today + timedelta(days=15), time(11, 30), "Ana Torres",
         "Contrato indefinido (periodo de prueba)", "", "Media", ""),
        (today + timedelta(days=25), None, "Luis Fernández",
         "Contrato temporal", "", "Media", ""),
        (today - timedelta(days=1), None, "Pedro Sánchez",
         "Contrato temporal", "¡Vencido ayer! Regularizar", "Crítica", ""),
        (today, time(16, 0), "Lucía Márquez",
         "Contrato de sustitución", "Último día", "Alta", ""),
        (today + timedelta(days=45), None, "Elena Díaz",
         "Contrato indefinido", "", "Baja", ""),
    ]


def write(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tareas"

    sheet.append(HEADERS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2C3E50")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in build_rows():
        sheet.append(list(row))

    # Formato de fecha/hora y anchos de columna.
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row=row_index, column=1).number_format = "DD/MM/YYYY"
        cell = sheet.cell(row=row_index, column=2)
        if cell.value is not None:
            cell.number_format = "HH:MM"

    widths = {"A": 12, "B": 8, "C": 22, "D": 34, "E": 30, "F": 12, "G": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    workbook.save(path)
    return path


def main() -> int:
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("tareas_prueba.xlsx")
    write(target.resolve())
    print(f"Excel de pruebas creado en: {target.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
